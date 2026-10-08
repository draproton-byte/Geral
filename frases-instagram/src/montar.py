"""Junta as frases dos 3 cursos, grava data/frases.json (entrada das artes) e a planilha única (xlsx + csv)."""
import csv, json, re, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

S = '/tmp/claude-0/-home-user-Geral/a8e795fc-5cc3-55e0-a62d-37078afc47c4/scratchpad/'
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = 'https://github.com/draproton-byte/Geral/blob/claude/ecstatic-meitner-ffcqdp/frases-instagram/artes/'

def load(n): return json.load(open(S + n, encoding='utf8'))

clube = [f for f in load('frases-clube-secreto.json') if int(f['id'][2:]) <= 20]  # CS21-30 são do livro Limite Zero (Joe Vitale)
desafio = load('frases-desafio.json')
im_new = load('frases-imersao-novas.json') + load('frases-imersao-lacunas.json')
old = json.load(open(os.path.join(ROOT, '..', 'imersao-dpm', 'src', 'frases.json'), encoding='utf8'))
im_old = [{'id': 'IM' + f['id'][1:], 'texto': f['texto'], 'ref': f['ref']} for f in old]

CURSOS = [('clube', 'Clube Secreto', clube), ('desafio', 'Desafio', desafio), ('imersao', 'Imersão', im_old + im_new)]
arts = []
for key, nome, lst in CURSOS:
    for f in lst:
        f.update(curso_key=key, curso=nome)
        arts.append(f)
json.dump(arts, open(os.path.join(ROOT, 'data', 'frases.json'), 'w', encoding='utf8'), ensure_ascii=False, indent=1)

def links(f):
    d = {'clube': 'clube', 'desafio': 'desafio', 'imersao': 'imersao'}[f['curso_key']]
    return (BASE + f"{d}/identidade-produto/{f['id']}.jpg", BASE + f"{d}/estilo-referencia/{f['id']}.jpg")

HEAD = ['ID', 'Curso', 'Formato', 'Tema', 'Referência da aula', 'Criativo pronto', 'Gancho', 'Copy / Roteiro',
        'Argumentação', 'Proposta', 'Estratégia', 'CTA', 'Link da arte · Identidade do produto', 'Link da arte · Estilo Dra. (pergaminho)', 'Status']
rows = []

# 1) Planejamento original da Imersão (cards, carrosséis, reels, caixinhas), com links das artes que já existem
cobertos = set()
for r in csv.DictReader(open(os.path.join(ROOT, '..', 'imersao-dpm', 'planilha', 'planejamento-conteudo-imersao-dpm.comma.csv'), encoding='utf8')):
    m = re.findall(r'F(\d\d)', r['Criativo pronto'])
    li = le = ''
    if r['Formato'] == 'Card' and m:
        li, le = links({'curso_key': 'imersao', 'id': 'IM' + m[0]})
        cobertos.update('IM' + x for x in m)
    rows.append([r['ID'], 'Imersão', r['Formato'], r['Tema'], r['Referência da aula'], r['Criativo pronto'], r['Gancho'],
                 r['Copy / Roteiro'], r['Argumentação'], r['Proposta'], r['Estratégia'], r['CTA'], li, le, r['Status']])

# 2) Frases novas (cards) dos 3 cursos; frases antigas da Imersão que já têm card planejado não repetem
for key, nome, lst in CURSOS:
    for f in lst:
        if f['id'] in cobertos: continue
        li, le = links(f)
        rows.append([f['id'], nome, 'Card (frase)', f['tema'] if 'tema' in f else f['ref'].split('·')[-1].strip(), f['ref'],
                     'Arte pronta · 2 estilos (feed 1080x1440)', f['texto'].replace('[', '').replace(']', ''),
                     f.get('legenda', 'A produzir legenda'), f.get('argumentacao', ''), f.get('proposta', ''),
                     f.get('estrategia', ''), f.get('cta', ''), li, le,
                     'Arte pronta · aguardando aprovação' if 'legenda' in f else 'Arte pronta · legenda a produzir'])

order = {'Clube Secreto': 0, 'Desafio': 1, 'Imersão': 2}
rows.sort(key=lambda r: (order[r[1]], r[0]))
json.dump(rows, open(os.path.join(ROOT, 'data', 'planilha.json'), 'w', encoding='utf8'), ensure_ascii=False)

os.makedirs(os.path.join(ROOT, 'planilha'), exist_ok=True)
with open(os.path.join(ROOT, 'planilha', 'planilha-frases-dra-proton.csv'), 'w', encoding='utf-8-sig', newline='') as fh:
    w = csv.writer(fh); w.writerow(HEAD); w.writerows(rows)

wb = Workbook(); ws = wb.active; ws.title = 'Conteúdo'
ws.append(HEAD)
for r in rows: ws.append(r)
hf = PatternFill('solid', fgColor='1F2A44')
for c in ws[1]:
    c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10); c.fill = hf
    c.alignment = Alignment(wrap_text=True, vertical='center')
thin = Side(style='thin', color='D0D0D0')
cor = {'Clube Secreto': 'F7EFD9', 'Desafio': 'EFE3F2', 'Imersão': 'E3EAF6'}
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name='Arial', size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = Border(bottom=thin)
    row[1].fill = PatternFill('solid', fgColor=cor[row[1].value])
    for i in (12, 13):
        if row[i].value:
            row[i].hyperlink = row[i].value; row[i].value = 'abrir arte'; row[i].font = Font(name='Arial', size=10, color='0563C1', underline='single')
for i, wd in enumerate([8, 14, 16, 24, 40, 26, 40, 60, 40, 40, 36, 32, 20, 22, 26], 1):
    ws.column_dimensions[get_column_letter(i)].width = wd
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 32
wb.save(os.path.join(ROOT, 'planilha', 'planilha-frases-dra-proton.xlsx'))
from collections import Counter
print(len(rows), Counter((r[1], r[2]) for r in rows))
